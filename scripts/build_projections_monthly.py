"""Rebuild the Projections tab with 12 monthly columns grouped under each year column.

Layout:  A labels | B At Close | C:N  Y1 months | O  Year 1 | P:AA Y2 months | AB Year 2 | ...
Month columns are outline-grouped (level 1) and hidden; the year column is the visible summary.
"""
import copy
import re

import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, column_index_from_string

SRC = "Smart_VAs_Acquisition_Model_2.xlsx"
DST = "Smart_VAs_Acquisition_Model_2_Monthly.xlsx"

OLD_YEAR_COLS = ["C", "D", "E", "F", "G"]          # Year 1 .. Year 5 today
YEAR_COL_IDX = [15 + 13 * y for y in range(5)]      # O, AB, AO, BB, BO
YEAR_COL = [get_column_letter(i) for i in YEAR_COL_IDX]
MONTH_COL_IDX = [[2 + 13 * y + m for m in range(1, 13)] for y in range(5)]

COLMAP = {"B": "B"}
COLMAP.update({old: new for old, new in zip(OLD_YEAR_COLS, YEAR_COL)})

# ---------------------------------------------------------------- formula tools
REF = re.compile(
    r"(?P<sheet>(?:'[^']*'|[A-Za-z_][A-Za-z0-9_.]*)!)?"
    r"(?P<abscol>\$?)(?P<col>[A-Z]{1,3})(?P<absrow>\$?)(?P<row>\d+)"
)


def _walk(formula, fn):
    """Apply fn(col_letter) -> new_col_letter to every un-qualified local reference."""
    def sub(m):
        if m.group("sheet"):
            return m.group(0)
        new = fn(m.group("col"))
        return f"{m.group('abscol')}{new}{m.group('absrow')}{m.group('row')}"

    return REF.sub(sub, formula)


def to_year_cols(formula):
    """Remap a Projections formula written for the old C:G grid onto the new year columns."""
    return _walk(formula, lambda c: COLMAP.get(c, c))


def shift(formula, delta):
    """Shift local references by `delta` columns (used to clone a Year-1 formula into a month)."""
    return _walk(formula, lambda c: get_column_letter(column_index_from_string(c) + delta))


# ---------------------------------------------------------------- row behaviour
HEADER_ROWS = [6, 44, 66, 96, 109, 122, 130, 138, 148]

# Annual flow allocated evenly across the twelve months of its year.
DIV12 = {7, 14, 15, 16, 17, 20, 21, 27, 31, 32, 39, 74, 78, 79, 80, 81,
         113, 118, 124, 125, 132, 133}

# Rate / ratio assumptions that are constant within the year.
SAME_AS_YEAR = {8, 142}

# Point-in-time balances: straight-line between the prior year-end and this year-end.
INTERP = {46, 47, 49, 50, 54, 55, 59, 60, 111, 116}

# Opening-balance rows that chain off the immediately preceding column.
# row -> (source row in the previous column, formula used for Year 1 Month 1)
CHAIN = {
    88: (46, None),                        # Beginning Cash <- prior column's Cash
    123: (127, "=Assumptions!E32"),        # SBA loan opening balance
    131: (135, "=Assumptions!E17"),        # Seller note opening balance
}

# Period movements derived from the change in a balance-sheet row.
# row -> (balance row, orientation).  "use" = -(closing - opening), the sign convention for
# a working-capital movement; "source" = -(opening - closing), the convention the model uses
# for debt principal repayment (an outflow).
DELTA = {70: (47, "use"), 82: (54, "source"), 83: (55, "source")}

# Valuation rows that stay annual (discounting is applied on an annual basis).
ANNUAL_ONLY = {157, 158, 159, 161, 162, 163, 164, 166, 167, 169, 170}

NOTE = ("Months are an even 1/12 allocation of each annual figure (balances are "
        "straight-lined between year-ends), so every month group foots exactly to its "
        "year column and no annual result changes.")


def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb["Projections"]

    # ---- snapshot the existing grid before anything is overwritten
    old = {}
    for r in range(1, ws.max_row + 1):
        for col in ["B"] + OLD_YEAR_COLS:
            c = ws[f"{col}{r}"]
            old[(col, r)] = (c.value, c._style)

    max_row = ws.max_row

    # ---- clear the old year grid so nothing survives where a month column lands
    for old_col in OLD_YEAR_COLS:
        for r in range(1, max_row + 1):
            ws[f"{old_col}{r}"].value = None

    # ---- year columns: original formulas, references remapped to the new grid
    for old_col, new_col in zip(OLD_YEAR_COLS, YEAR_COL):
        for r in range(1, max_row + 1):
            val, style = old[(old_col, r)]
            if val is None:
                continue
            if isinstance(val, str) and val.startswith("="):
                val = to_year_cols(val)
            tgt = ws[f"{new_col}{r}"]
            tgt.value = val
            tgt._style = copy.copy(style)

    # Sum of PV picks up the five year columns explicitly (the old SUM(C159:G159)
    # would otherwise sweep in every monthly column).
    ws[f"{YEAR_COL[0]}161"] = "=" + "+".join(f"{c}159" for c in YEAR_COL)

    # ---- monthly columns
    for y in range(5):
        for m, cidx in enumerate(MONTH_COL_IDX[y], start=1):
            col = get_column_letter(cidx)
            ycol = YEAR_COL[y]
            prev_col = get_column_letter(cidx - 1)   # Y1M1 -> B, Y2M1 -> Year 1 col, ...
            prev_year_col = "B" if y == 0 else YEAR_COL[y - 1]
            delta = cidx - column_index_from_string("C")  # clone Year-1 formulas across

            for r in range(1, max_row + 1):
                src_val, src_style = old[("C", r)]
                cell = ws[f"{col}{r}"]

                if r in HEADER_ROWS:
                    cell.value = f"M{m}"
                    cell._style = copy.copy(src_style)
                    continue
                if src_val is None:
                    continue
                if r in ANNUAL_ONLY:
                    # Valuation rows stay annual; keep the month cells blank but styled.
                    cell._style = copy.copy(src_style)
                    continue

                if r in DIV12:
                    cell.value = f"=${ycol}{r}/12"
                elif r in SAME_AS_YEAR:
                    cell.value = f"=${ycol}{r}"
                elif r in INTERP:
                    cell.value = (f"=${prev_year_col}{r}+(${ycol}{r}-${prev_year_col}{r})*{m}/12")
                elif r in CHAIN:
                    src_row, first = CHAIN[r]
                    cell.value = first if (y == 0 and m == 1 and first) else f"={prev_col}{src_row}"
                elif r in DELTA:
                    br, orient = DELTA[r]
                    cell.value = (f"=-({col}{br}-{prev_col}{br})" if orient == "use"
                                  else f"=-({prev_col}{br}-{col}{br})")
                elif isinstance(src_val, str) and src_val.startswith("="):
                    cell.value = shift(src_val, delta)
                else:
                    cell.value = src_val

                cell._style = copy.copy(src_style)

    # ---- outline grouping: months hidden, year column is the visible summary
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.sheet_properties.outlinePr.showOutlineSymbols = True
    # The original file carried one <col> entry spanning B:G; the new grid sets every column
    # explicitly, so narrow it to B to avoid overlapping range definitions.
    ws.column_dimensions["B"].min = ws.column_dimensions["B"].max = 2
    for y in range(5):
        for cidx in MONTH_COL_IDX[y]:
            cd = ws.column_dimensions[get_column_letter(cidx)]
            cd.width = 12.0
            cd.outlineLevel = 1
            cd.hidden = True
            cd.collapsed = False
        yd = ws.column_dimensions[YEAR_COL[y]]
        yd.width = 15.0
        yd.outlineLevel = 0
        yd.hidden = False
        yd.collapsed = True

    # ---- document the allocation where the reader will see it
    ws["A3"] = (ws["A3"].value or "") + (
        " Monthly detail: each year is grouped into 12 monthly columns (M1–M12), hidden by "
        "default — click the [+] above the year column to expand. " + NOTE)
    ws["O6"].comment = Comment(NOTE, "Acquisition Model")
    ws["A157"].comment = Comment(
        "Discounting is applied on an annual basis; rows 157-159 are shown in the year "
        "columns only.", "Acquisition Model")

    # ---- Analysis tab points at the relocated year columns
    an = wb["Analysis"]
    pat = re.compile(r"(Projections!)(\$?)([A-G])(\$?\d+)")
    for row in an.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and "Projections!" in c.value:
                c.value = pat.sub(lambda m: f"{m.group(1)}{m.group(2)}{COLMAP[m.group(3)]}{m.group(4)}",
                                  c.value)

    wb.save(DST)
    print("wrote", DST)


if __name__ == "__main__":
    main()
